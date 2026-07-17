import docx
from docxtpl import DocxTemplate

doc = docx.Document()
table = doc.add_table(rows=2, cols=2)
table.cell(0, 0).text = "Name"
table.cell(0, 1).text = "Value"

table.cell(1, 0).text = "{% tr for item in items %}{{ item.name }}"
table.cell(1, 1).text = "{{ item.val }}{% tr endfor %}"

doc.save("test_table.docx")

tpl = DocxTemplate("test_table.docx")
context = {'items': [{'name': 'A', 'val': 1}, {'name': 'B', 'val': 2}]}
tpl.render(context)
tpl.save("test_table_out.docx")
print("Render successful.")
